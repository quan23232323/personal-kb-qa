"""文档处理流水线：上传 → 解析 → 分块 → 向量化 → 入库。"""
import io
import traceback
import uuid
from pathlib import Path

from config import config
from database import db
from utils import file_utils
from utils.time_utils import now_iso
from . import embedding_service, retrieval_service


def _upload_dir(kb_id: str) -> Path:
    d = config.resolve_path(config.get("storage.upload_dir", "./data/uploads")) / kb_id
    d.mkdir(parents=True, exist_ok=True)
    return d


def get_document(doc_id: str) -> dict | None:
    with db() as conn:
        row = conn.execute("SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
    return dict(row) if row else None


def list_documents(kb_id: str) -> list[dict]:
    with db() as conn:
        rows = conn.execute(
            "SELECT * FROM documents WHERE kb_id = ? ORDER BY created_at DESC", (kb_id,)
        ).fetchall()
    return [dict(r) for r in rows]


def create_document_from_upload(
    kb_id: str, original_filename: str, file_type: str, content: bytes
) -> dict:
    """保存原始文件并创建 processing 状态的 Document 记录。"""
    safe_name = file_utils.sanitize_filename(original_filename)
    stored_name = f"{uuid.uuid4().hex}_{safe_name}"
    path = _upload_dir(kb_id) / stored_name
    path.write_bytes(content)

    doc_id = str(uuid.uuid4())
    now = now_iso()
    with db() as conn:
        conn.execute(
            "INSERT INTO documents(id, kb_id, filename, file_path, file_type, file_size, status, created_at) "
            "VALUES (?,?,?,?,?,?,?,?)",
            (doc_id, kb_id, safe_name, str(path), file_type, len(content), "processing", now),
        )
    return get_document(doc_id)


def _read_text_file(path: str) -> str:
    """读 txt/md：依次尝试 utf-8-sig → gbk（Windows 中文常见），失败再宽松解码。"""
    data = Path(path).read_bytes()
    for enc in ("utf-8-sig", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _table_rows_to_text(rows: list[list[str]], sheet_name: str | None) -> str:
    """把表格行序列化为可检索文本：首行作为表头，逐行「表头: 值」拼接。

    「表头: 值」形式比裸表格更利于语义检索（每个单元格都带上字段含义）。
    """
    if not rows:
        return ""
    header = [str(h or "").strip() for h in rows[0]]
    lines = [f"【{sheet_name}】"] if sheet_name else []
    for row in rows[1:]:
        cells = [str(c or "").strip() for c in row]
        if not any(cells):
            continue
        pairs = []
        for i, cell in enumerate(cells):
            if not cell:
                continue
            if i < len(header) and header[i]:
                pairs.append(f"{header[i]}: {cell}")
            else:
                pairs.append(cell)
        lines.append("；".join(pairs))
    return "\n".join(lines)


def _extract_csv(path: str) -> str:
    """读 csv：中文 Excel 导出的 csv 常为 GBK/带 BOM，做编码回退。"""
    import csv

    data = Path(path).read_bytes()
    text = None
    for enc in ("utf-8-sig", "gbk"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode("utf-8", errors="ignore")
    rows = list(csv.reader(io.StringIO(text)))
    return _table_rows_to_text(rows, None)


def _extract_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        sheet_text = _table_rows_to_text(rows, ws.title)
        if sheet_text:
            parts.append(sheet_text)
    wb.close()
    return "\n\n".join(parts)


def _extract_xls(path: str) -> str:
    import xlrd

    book = xlrd.open_workbook(path)
    parts = []
    for ws in book.sheets():
        rows = [ws.row_values(r) for r in range(ws.nrows)]
        sheet_text = _table_rows_to_text(rows, ws.name)
        if sheet_text:
            parts.append(sheet_text)
    return "\n\n".join(parts)


def _extract_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                t = _table_rows_to_text(rows, None)
                if t:
                    texts.append(t)
        if slide.has_notes_slide:
            t = (slide.notes_slide.notes_text_frame.text or "").strip()
            if t:
                texts.append(t)
        if texts:
            parts.append(f"【第{i}页】\n" + "\n".join(texts))
    return "\n\n".join(parts)


def _extract_text(file_path: str, file_type: str) -> str:
    if file_type in ("txt", "md"):
        return _read_text_file(file_path)
    if file_type == "csv":
        return _extract_csv(file_path)
    if file_type == "xlsx":
        return _extract_xlsx(file_path)
    if file_type == "xls":
        return _extract_xls(file_path)
    if file_type == "pptx":
        return _extract_pptx(file_path)
    if file_type == "pdf":
        from langchain_community.document_loaders import PyPDFLoader

        docs = PyPDFLoader(file_path).load()
        return "\n\n".join(d.page_content for d in docs)
    if file_type == "docx":
        from langchain_community.document_loaders import Docx2txtLoader

        docs = Docx2txtLoader(file_path).load()
        return "\n\n".join(d.page_content for d in docs)
    # 旧版二进制格式（doc/ppt）没有可靠的纯 Python 解析方案，明确告知转换
    raise ValueError(
        f"不支持的文件类型: {file_type}"
        + ("（旧版二进制格式请另存为 docx/pptx 后上传）" if file_type in ("doc", "ppt") else "")
    )


def _split_text(text: str) -> list[str]:
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.get("document.chunk_size", 500),
        chunk_overlap=config.get("document.chunk_overlap", 50),
        separators=["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""],
    )
    return splitter.split_text(text)


def process_document(doc_id: str, kb_id: str, file_path: str, file_type: str, filename: str) -> None:
    """后台流水线：解析 → 分块 → 向量化 → 入 ChromaDB → 更新状态。"""
    try:
        text = _extract_text(file_path, file_type)
        chunks = [c for c in _split_text(text) if c.strip()]
        if not chunks:
            raise ValueError("文档无有效文本内容")
        embeddings = embedding_service.embed(chunks)
        ids = [f"{doc_id}_{i}" for i in range(len(chunks))]
        retrieval_service.add_chunks(kb_id, ids, chunks, embeddings, doc_id, filename)
        with db() as conn:
            conn.execute(
                "UPDATE documents SET status='completed', chunk_count=?, error_message=NULL WHERE id=?",
                (len(chunks), doc_id),
            )
    except Exception as e:
        with db() as conn:
            conn.execute(
                "UPDATE documents SET status='failed', error_message=? WHERE id=?",
                (f"{e}\n{traceback.format_exc()}", doc_id),
            )


def delete_document(doc_id: str) -> bool:
    doc = get_document(doc_id)
    if not doc:
        return False
    retrieval_service.delete_document_vectors(doc["kb_id"], doc_id)
    with db() as conn:
        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
    try:
        Path(doc["file_path"]).unlink(missing_ok=True)
    except Exception:
        pass
    return True


def start_reprocess(doc_id: str) -> None:
    """重新处理前：清旧向量 + 置 processing 状态。"""
    doc = get_document(doc_id)
    if not doc:
        raise ValueError("文档不存在")
    retrieval_service.delete_document_vectors(doc["kb_id"], doc_id)
    with db() as conn:
        conn.execute("UPDATE documents SET status='processing', error_message=NULL WHERE id=?", (doc_id,))


def reset_stale_processing() -> int:
    """启动时把上次进程中断遗留的 processing 文档标记为 failed（抗 --reload / 崩溃）。"""
    with db() as conn:
        cur = conn.execute(
            "UPDATE documents SET status='failed', error_message=? WHERE status='processing'",
            ("处理被中断，请重新处理",),
        )
        return cur.rowcount
