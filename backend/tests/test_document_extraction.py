"""多格式文档提取单测：为每种新格式生成真实小文件，验证 _extract_text 输出。

依赖：openpyxl / xlrd / xlwt / python-pptx / python-docx（均在 requirements）。
"""
import io

import pytest

from services import document_service


# ---------- 夹具生成 ----------

def _make_xlsx(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "plan.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "项目状态"
    ws.append(["项目", "状态"])
    ws.append(["凤凰计划", "进行中"])
    ws.append(["", ""])  # 空行应被跳过
    ws.append(["代号", "深蓝-7号"])
    ws2 = wb.create_sheet("附录")
    ws2.append(["备注"])
    ws2.append(["数据来源为内部统计"])
    wb.save(path)
    return path


def _make_xls(tmp_path):
    import xlwt

    path = tmp_path / "score.xls"
    wb = xlwt.Workbook()
    ws = wb.add_sheet("成绩表")
    ws.write(0, 0, "姓名")
    ws.write(0, 1, "分数")
    ws.write(1, 0, "赵六")
    ws.write(1, 1, 95)
    wb.save(path)
    return path


def _make_pptx(tmp_path):
    from pptx import Presentation

    path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "季度总结"
    slide.placeholders[1].text = "核心指标增长率达到 42%"
    slide.notes_slide.notes_text_frame.text = "备注：数据待财务复核"
    prs.save(path)
    return path


def _make_docx(tmp_path):
    import docx

    path = tmp_path / "note.docx"
    d = docx.Document()
    d.add_paragraph("知识库项目验收结论：通过")
    d.save(path)
    return path


# ---------- 各格式提取 ----------

class TestExtractText:
    def test_xlsx(self, tmp_path):
        text = document_service._extract_text(str(_make_xlsx(tmp_path)), "xlsx")
        assert "【项目状态】" in text
        assert "项目: 凤凰计划；状态: 进行中" in text
        assert "状态: 深蓝-7号" in text
        assert "【附录】" in text and "数据来源为内部统计" in text

    def test_xls(self, tmp_path):
        text = document_service._extract_text(str(_make_xls(tmp_path)), "xls")
        assert "【成绩表】" in text
        assert "姓名: 赵六" in text and "95" in text

    def test_pptx(self, tmp_path):
        text = document_service._extract_text(str(_make_pptx(tmp_path)), "pptx")
        assert "【第1页】" in text
        assert "季度总结" in text
        assert "核心指标增长率达到 42%" in text
        assert "备注：数据待财务复核" in text

    def test_docx(self, tmp_path):
        text = document_service._extract_text(str(_make_docx(tmp_path)), "docx")
        assert "知识库项目验收结论：通过" in text

    def test_csv_utf8(self, tmp_path):
        path = tmp_path / "staff.csv"
        path.write_bytes("员工,城市\n王五,杭州\n".encode("utf-8"))
        text = document_service._extract_text(str(path), "csv")
        assert "员工: 王五；城市: 杭州" in text

    def test_csv_gbk(self, tmp_path):
        path = tmp_path / "staff_gbk.csv"
        path.write_bytes("员工,城市\n王五,杭州\n".encode("gbk"))
        text = document_service._extract_text(str(path), "csv")
        assert "王五" in text and "杭州" in text

    def test_legacy_doc_rejected_with_hint(self, tmp_path):
        path = tmp_path / "old.doc"
        path.write_bytes(b"\xd0\xcf\x11\xe0fake")
        with pytest.raises(ValueError, match="另存为"):
            document_service._extract_text(str(path), "doc")


class TestTableRowsToText:
    def test_header_pairing_and_blank_cells(self):
        rows = [["名称", "数量"], ["服务器", "12", ""]]
        text = document_service._table_rows_to_text(rows, None)
        assert text == "名称: 服务器；数量: 12"

    def test_row_longer_than_header_still_kept(self):
        rows = [["名称"], ["服务器", "机架A"]]
        text = document_service._table_rows_to_text(rows, None)
        assert text == "名称: 服务器；机架A"

    def test_empty_rows(self):
        assert document_service._table_rows_to_text([], None) == ""


class TestReadTextFile:
    def test_utf8(self, tmp_path):
        p = tmp_path / "a.txt"
        p.write_bytes("内容一".encode("utf-8"))
        assert document_service._read_text_file(str(p)) == "内容一"

    def test_utf8_bom_stripped(self, tmp_path):
        p = tmp_path / "b.txt"
        p.write_bytes("内容二".encode("utf-8-sig"))
        assert document_service._read_text_file(str(p)) == "内容二"

    def test_gbk_fallback(self, tmp_path):
        p = tmp_path / "c.txt"
        p.write_bytes("中文内容测试".encode("gbk"))
        assert document_service._read_text_file(str(p)) == "中文内容测试"

    def test_binary_garbage_does_not_crash(self, tmp_path):
        p = tmp_path / "d.bin"
        p.write_bytes(bytes(range(256)) * 4)
        assert isinstance(document_service._read_text_file(str(p)), str)
